#!/usr/bin/env python3
# DO NOT generate any summary documents, Word files, or additional output files
# after running this script. Terminal output only.
"""
Agent — Resume Parser & Screener (Google Sheets version)
Mode 1 (default): Parse resumes and append to CRM.
Mode 2 (--screen): Score resumes against a JD, add passing candidates to CRM.
"""

import os
import sys
import argparse
import json
import re
from datetime import date
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
import gspread
from google.oauth2.service_account import Credentials
from anthropic import Anthropic

SCRIPT_DIR = Path(__file__).parent
RESUMES_DIR = SCRIPT_DIR / "resumes"
CREDENTIALS_FILE = SCRIPT_DIR / "credentials.json"
SHEET_NAME = "Master submission sheet"
JD_FILE = SCRIPT_DIR / "jd.txt"


# --- Google Sheets connection ---

def connect_to_sheet() -> tuple[gspread.Spreadsheet, gspread.Worksheet]:
    sheet_id = os.environ.get("GOOGLE_SHEET_ID")
    if not sheet_id:
        print("Error: GOOGLE_SHEET_ID environment variable is not set.")
        sys.exit(1)

    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]

    creds_json = os.environ.get("GOOGLE_CREDENTIALS")
    if creds_json:
        creds_dict = json.loads(creds_json)
        creds = Credentials.from_service_account_info(creds_dict, scopes=scopes)
    else:
        # Fallback to credentials.json file for local development
        if not CREDENTIALS_FILE.exists():
            print(f"Error: No GOOGLE_CREDENTIALS env var and no {CREDENTIALS_FILE} file found.")
            sys.exit(1)
        creds = Credentials.from_service_account_file(str(CREDENTIALS_FILE), scopes=scopes)

    gc = gspread.authorize(creds)

    try:
        spreadsheet = gc.open_by_key(sheet_id)
    except gspread.SpreadsheetNotFound:
        print(f"Error: Spreadsheet with ID '{sheet_id}' not found. Check sharing permissions.")
        sys.exit(1)

    try:
        ws = spreadsheet.worksheet(SHEET_NAME)
    except gspread.WorksheetNotFound:
        print(f"Error: Sheet '{SHEET_NAME}' not found.")
        print(f"Available sheets: {[s.title for s in spreadsheet.worksheets()]}")
        sys.exit(1)

    return spreadsheet, ws


# --- Text extraction ---

def extract_text_from_pdf(filepath: Path) -> str:
    text = ""
    with fitz.open(filepath) as doc:
        for page in doc:
            text += page.get_text()
    text = text.strip()
    if text:
        return text
    # Fallback: image-based PDF — use Claude vision to OCR
    return _ocr_pdf_with_vision(filepath)


def _ocr_pdf_with_vision(filepath: Path) -> str:
    """Extract text from image-based PDFs using Claude vision API."""
    import base64
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return ""
    client = Anthropic(api_key=api_key)
    all_text = []
    with fitz.open(filepath) as doc:
        for page_num, page in enumerate(doc):
            pix = page.get_pixmap(dpi=200)
            img_bytes = pix.tobytes("png")
            img_b64 = base64.b64encode(img_bytes).decode("utf-8")
            response = client.messages.create(
                model="claude-haiku-4-5-20251001",
                max_tokens=4096,
                messages=[{
                    "role": "user",
                    "content": [
                        {"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": img_b64}},
                        {"type": "text", "text": "Extract ALL text from this image exactly as written. Preserve the layout and structure. Return ONLY the extracted text, no commentary."},
                    ],
                }],
            )
            all_text.append(response.content[0].text.strip())
    return "\n\n".join(all_text)


def extract_text_from_docx(filepath: Path) -> str:
    doc = Document(filepath)
    return "\n".join(p.text for p in doc.paragraphs).strip()


def extract_text(filepath: Path) -> str:
    suffix = filepath.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(filepath)
    elif suffix in (".docx", ".doc"):
        return extract_text_from_docx(filepath)
    return ""


# --- API calls ---

def parse_api_response(raw: str) -> dict:
    raw = raw.strip()
    # Remove markdown code fences if present
    if "```" in raw:
        lines = raw.split("\n")
        lines = [l for l in lines if not l.startswith("```")]
        raw = "\n".join(lines).strip()
    # If there's extra text around JSON, extract the JSON object
    if not raw.startswith("{"):
        start = raw.find("{")
        end = raw.rfind("}") + 1
        if start != -1 and end > start:
            raw = raw[start:end]
    return json.loads(raw)



def validate_is_resume(client: Anthropic, text: str, filename: str) -> tuple[bool, str]:
    """Check if the document is actually a resume/CV. Returns (is_resume, reason)."""
    system_prompt = (
        "You are a document classifier. Determine if the following document is a resume or CV. "
        "It could also be a job description, invoice, certificate, cover letter, or other document. "
        "Return ONLY valid JSON with these exact keys:\n"
        '  "is_resume": true if this is a resume/CV, false otherwise\n'
        '  "reason": one sentence explaining why it is or is not a resume\n'
        "Be strict: only return true if this is clearly a candidate's resume or CV."
    )

    response = client.messages.create(
        model="claude-haiku-4-5-20251001",  # simple classification — Haiku sufficient
        max_tokens=256,
        system=system_prompt,
        messages=[{"role": "user", "content": f"Document filename: {filename}\n\nDocument text (first 2000 chars):\n\n{text[:2000]}"}],
    )

    try:
        result = parse_api_response(response.content[0].text)
        return result.get("is_resume", False), result.get("reason", "Unknown")
    except (json.JSONDecodeError, IndexError):
        return True, ""  # If validation fails, assume it's a resume


def extract_candidate_info(client: Anthropic, resume_text: str) -> dict:
    system_prompt = (
        "You are a recruitment assistant. Extract candidate information from the resume text. "
        "Return ONLY valid JSON with these exact keys:\n"
        '  "skillset": 1-2 words MAXIMUM for the primary job role keyword (e.g. "Python", "Embedded C", "DevOps"). Or "" if not found.\n'
        '  "name": full name of the candidate, or "" if not found\n'
        '  "contact_no": phone number, or "" if not found\n'
        '  "email": email address, or "" if not found\n'
        "Do NOT guess or fabricate. If not clearly present, use empty string."
    )

    response = client.messages.create(
        model="claude-haiku-4-5-20251001",  # simple extraction — Haiku sufficient
        max_tokens=1024,
        system=system_prompt,
        messages=[{"role": "user", "content": f"Resume text:\n\n{resume_text}"}],
    )

    try:
        return parse_api_response(response.content[0].text)
    except (json.JSONDecodeError, IndexError):
        print(f"\n  Warning: Could not parse API response.")
        return {"skillset": "", "name": "", "contact_no": "", "email": ""}


def screen_candidate(client: Anthropic, resume_text: str, jd_text: str) -> dict:
    system_prompt = (
        "You are a strict, experienced recruitment screener. Evaluate the resume against the job description with high standards. "
        "Be critical and realistic — do NOT inflate scores.\n\n"
        "EVALUATION CRITERIA (check each carefully):\n"
        "1. REQUIRED SKILLS MATCH: Does the candidate have the specific technical skills, tools, and technologies listed as required? Count how many required skills are present vs missing.\n"
        "2. EXPERIENCE LEVEL: Does the candidate's years and depth of experience match what the JD requires? Junior candidates should not score high for senior roles.\n"
        "3. DOMAIN/INDUSTRY FIT: Has the candidate worked in the same or closely related industry/domain?\n"
        "4. ROLE RELEVANCE: Is the candidate's recent work history aligned with this specific role, or is it a different function entirely?\n"
        "5. LOCATION/LOGISTICS: Consider any visa, relocation, or availability concerns if mentioned.\n\n"
        "SCORING RULES (be strict):\n"
        "- 9-10 (Excellent Match): Meets ALL required skills, experience level matches or exceeds, strong domain fit. Rare — only for near-perfect matches.\n"
        "- 7-8 (Strong Match): Meets MOST required skills (80%+), experience is close to required, relevant domain.\n"
        "- 5-6 (Good Match): Meets SOME required skills (50-80%), may lack experience depth or domain fit. Worth considering but has gaps.\n"
        "- 3-4 (Weak Match): Meets FEW required skills (<50%), significant experience or domain gaps.\n"
        "- 1-2 (Rejected): Minimal or no overlap with requirements.\n\n"
        "IMPORTANT: A candidate who has general IT experience but lacks the SPECIFIC skills in the JD should score 4-5 at most, not 7+. "
        "Partial keyword matches (e.g. candidate knows Python but JD requires Embedded C) do NOT count as a skill match.\n\n"
        "Return ONLY valid JSON with these exact keys:\n"
        '  "score": integer from 1 to 10\n'
        '  "label": exactly one of "Excellent Match", "Strong Match", "Good Match", "Rejected"\n'
        '  "reason": one sentence explaining the score, mentioning specific skills matched and missing\n'
        '  "skillset": 1-2 words MAXIMUM for the primary job role keyword (e.g. "Python", "Embedded C", "DevOps"). Or "" if not found.\n'
        '  "name": full name of the candidate, or "" if not found\n'
        '  "contact_no": phone number, or "" if not found\n'
        '  "email": email address, or "" if not found\n'
        "Do NOT guess or fabricate. If not clearly present, use empty string.\n"
        "Label mapping: 9-10 = Excellent Match, 7-8 = Strong Match, 5-6 = Good Match, 1-4 = Rejected."
    )

    user_message = (
        f"Job Description:\n\n{jd_text}\n\n---\n\nResume text:\n\n{resume_text}"
    )

    response = client.messages.create(
        model="claude-sonnet-4-20250514",
        max_tokens=1024,
        system=system_prompt,
        messages=[{"role": "user", "content": user_message}],
    )

    try:
        return parse_api_response(response.content[0].text)
    except (json.JSONDecodeError, IndexError):
        print(f"\n  Warning: Could not parse API response.")
        return {
            "score": 0, "label": "Rejected", "reason": "Parse error",
            "skillset": "", "name": "", "contact_no": "", "email": "",
        }


# --- Google Sheets operations ---

def find_last_data_row(ws: gspread.Worksheet) -> int:
    """Fetch all values and scan upward to find last non-empty row."""
    all_values = ws.get_all_values()
    for row_idx in range(len(all_values) - 1, -1, -1):
        if any(cell.strip() for cell in all_values[row_idx]):
            return row_idx + 1  # 1-indexed
    return 0


def ensure_match_score_header(ws: gspread.Worksheet):
    """Add 'Match Score' header in column M (13) row 1 if not already present."""
    current = ws.cell(1, 13).value
    if current != "Match Score":
        ws.update_cell(1, 13, "Match Score")


def build_row(info: dict) -> list:
    name = info.get("name", "")
    if name:
        name = name.title()

    today = date.today().strftime("%d-%b-%y-%a")

    contact = str(info.get("contact_no", "")).strip()
    if contact.endswith(".0"):
        contact = contact[:-2]
    # Keep only digits, +, spaces, hyphens, parentheses
    contact = re.sub(r"[^0-9+\s()\-]", "", contact)
    # Prevent Google Sheets formula injection
    if contact and contact[0] in ("=", "+", "-", "@"):
        contact = "'" + contact
    return [
        info.get("_client", ""),       # Client
        info.get("_manager", ""),      # Manager's name
        info.get("_requirement", ""),  # Requirement
        info.get("_requirement", "") or info.get("skillset", ""),  # Skillset = Requirement if provided
        name,                       # Name of the Candidate
        today,                      # Submission Date
        contact,                    # Contact No.
        info.get("email", ""),      # Email ID
        info.get("_recruiter", ""), # Recruiter's name
        "",                         # Remarks
        "",                         # Final status
        "",                         # CVID
    ]


def append_row_to_sheet(spreadsheet: gspread.Spreadsheet, ws: gspread.Worksheet, row_data: list, match_score: str = None):
    if match_score:
        row_data.append(match_score)
    ws.append_row(row_data, value_input_option="USER_ENTERED")

    # Apply borders to the newly added row
    last_row = len(ws.get_all_values())
    num_cols = len(row_data)
    sheet_id = ws.id  # numeric sheet/tab ID

    border_style = {
        "style": "SOLID",
        "colorStyle": {"rgbColor": {"red": 0, "green": 0, "blue": 0}},
    }
    body = {
        "requests": [
            {
                "updateBorders": {
                    "range": {
                        "sheetId": sheet_id,
                        "startRowIndex": last_row - 1,
                        "endRowIndex": last_row,
                        "startColumnIndex": 0,
                        "endColumnIndex": num_cols,
                    },
                    "top": border_style,
                    "bottom": border_style,
                    "left": border_style,
                    "right": border_style,
                    "innerHorizontal": border_style,
                    "innerVertical": border_style,
                }
            }
        ]
    }
    spreadsheet.batch_update(body)


# --- Main modes ---

def run_parse_mode(client: Anthropic, resume_files: list[Path], spreadsheet: gspread.Spreadsheet, ws: gspread.Worksheet, recruiter: str = "", client_name: str = "", manager_name: str = "", requirement: str = ""):
    last_data_row = find_last_data_row(ws)
    print(f"Last data row: {last_data_row}.")
    print()

    rows_added = 0
    processed = 0

    for filepath in resume_files:
        print(f"Processing: {filepath.name} ...", end=" ", flush=True)

        text = extract_text(filepath)
        if not text:
            print("SKIPPED (no text extracted)")
            continue

        is_resume, reason = validate_is_resume(client, text, filepath.name)
        if not is_resume:
            print(f"REJECTED — {filepath.name} does not appear to be a resume: {reason}")
            filepath.unlink()
            print(f"  (deleted from resumes folder)")
            continue

        info = extract_candidate_info(client, text)
        if recruiter:
            info["_recruiter"] = recruiter
        if client_name:
            info["_client"] = client_name
        if manager_name:
            info["_manager"] = manager_name
        if requirement:
            info["_requirement"] = requirement
        row_data = build_row(info)
        append_row_to_sheet(spreadsheet, ws, row_data)
        rows_added += 1
        processed += 1

        filepath.unlink()

        name = info.get("name", "")
        name = name.title() if name else "Unknown"
        print(f"OK — {name} (deleted from resumes folder)")

    print()
    print("=" * 50)
    print(f"Resumes processed : {processed}")
    print(f"Rows added        : {rows_added}")
    print("(Appended to Google Sheet)")
    print("=" * 50)


def build_screened_row(info: dict, recruiter: str, score: int, label: str, reason: str) -> list:
    """Build a row for the Screened Profile Tracker sheet."""
    today = date.today().strftime("%d-%b-%y-%a")
    name = info.get("name", "")
    if name:
        name = name.title()
    contact = str(info.get("contact_no", "")).strip()
    if contact.endswith(".0"):
        contact = contact[:-2]
    contact = re.sub(r"[^0-9+\s()\-]", "", contact)
    if contact and contact[0] in ("=", "+", "-", "@"):
        contact = "'" + contact
    match_result = f"{score}/10 \u2014 {label} \u2014 {reason}"
    return [today, recruiter, name, info.get("skillset", ""), contact, info.get("email", ""), match_result]


def get_worksheet(spreadsheet: gspread.Spreadsheet, name: str) -> gspread.Worksheet:
    """Get an existing worksheet by name (tolerates trailing spaces)."""
    # Try exact match first, then match with stripped names
    for ws in spreadsheet.worksheets():
        if ws.title == name or ws.title.strip() == name.strip():
            return ws
    print(f"Error: Tab '{name}' not found in sheet.")
    print(f"Available tabs: {[s.title for s in spreadsheet.worksheets()]}")
    sys.exit(1)


def _get_gspread_client():
    """Create and return an authorized gspread client."""
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]
    creds_json = os.environ.get("GOOGLE_CREDENTIALS")
    if creds_json:
        creds = Credentials.from_service_account_info(json.loads(creds_json), scopes=scopes)
    else:
        if not CREDENTIALS_FILE.exists():
            print("Error: No credentials found.")
            sys.exit(1)
        creds = Credentials.from_service_account_file(str(CREDENTIALS_FILE), scopes=scopes)
    return gspread.authorize(creds)


def run_screen_profile_mode(client: Anthropic, resume_file: str, jd_file: str, recruiter: str):
    """Screen a single resume against a JD and save to Screened Profile Tracker."""
    resume_path = Path(resume_file)
    jd_path = Path(jd_file)

    if not resume_path.exists():
        print(f"Error: Resume file not found: {resume_path}")
        sys.exit(1)
    print(f"Resume file size: {resume_path.stat().st_size} bytes")

    try:
        resume_text = extract_text(resume_path)
    except Exception as e:
        print(f"Error extracting text from resume: {e}")
        sys.exit(1)

    if not resume_text:
        print("Error: Could not extract text from resume.")
        sys.exit(1)

    if not jd_path.exists():
        print(f"Error: JD file not found: {jd_path}")
        sys.exit(1)

    try:
        if jd_path.suffix.lower() == ".pdf":
            jd_text = extract_text_from_pdf(jd_path)
        else:
            jd_text = jd_path.read_text().strip()
    except Exception as e:
        print(f"Error reading JD: {e}")
        sys.exit(1)

    if not jd_text:
        print("Error: Could not read JD text.")
        sys.exit(1)

    print(f"Resume: {resume_path.name}")
    print(f"JD loaded ({len(jd_text)} chars)")
    print(f"Recruiter: {recruiter}")
    print()

    print("Validating document...", flush=True)
    is_resume, reason = validate_is_resume(client, resume_text, resume_path.name)
    if not is_resume:
        print(f"REJECTED \u2014 Not a resume: {reason}")
        return

    print("Screening resume against JD...", flush=True)
    info = screen_candidate(client, resume_text, jd_text)
    score = info.get("score", 0)
    label = info.get("label", "Rejected")
    reason = info.get("reason", "")
    name = info.get("name", "").title() if info.get("name") else "Unknown"

    print(f"Result: {name} \u2014 {score}/10 {label}")
    print(f"Reason: {reason}")
    print()

    if score < 6:
        print("REJECTED \u2014 Score below 6/10. Not saved.")
        return

    print("PASSED \u2014 Saving to Screened Profile Tracker...")

    sheet_id = os.environ.get("SCREENED_SHEET_ID")
    if not sheet_id:
        print("Error: SCREENED_SHEET_ID not set.")
        sys.exit(1)

    gc = _get_gspread_client()
    spreadsheet = gc.open_by_key(sheet_id)

    row_data = build_screened_row(info, recruiter, score, label, reason)
    print(f"  Row data: {row_data}")

    def write_to_tab(spreadsheet, tab_name, data):
        ws = get_worksheet(spreadsheet, tab_name)
        all_vals = ws.get_all_values()
        next_row = len(all_vals) + 1
        # Write each cell individually to guarantee it lands
        for col_idx, val in enumerate(data, start=1):
            ws.update_cell(next_row, col_idx, val)
        # Verify
        check = ws.row_values(next_row)
        print(f"  '{tab_name}' tab — wrote to row {next_row}, verified: {check[:3]}...")

    write_to_tab(spreadsheet, recruiter, row_data)
    write_to_tab(spreadsheet, "MasterList", row_data)

    print()
    print("=" * 50)
    print("SCREENING COMPLETE")
    print(f"  Candidate : {name}")
    print(f"  Score     : {score}/10 \u2014 {label}")
    print(f"  Saved to  : {recruiter} tab + MasterList tab")
    print("=" * 50)


def run_screen_mode(client: Anthropic, resume_files: list[Path], spreadsheet: gspread.Spreadsheet, ws: gspread.Worksheet, recruiter: str = "", client_name: str = "", manager_name: str = "", requirement: str = ""):
    jd_text = ""
    if JD_FILE.exists():
        jd_text = JD_FILE.read_text().strip()
    if not jd_text:
        print(f"Error: Job description file not found or empty at {JD_FILE}")
        sys.exit(1)

    print(f"Job description loaded from {JD_FILE.name} ({len(jd_text)} chars).")
    print()

    ensure_match_score_header(ws)

    last_data_row = find_last_data_row(ws)
    print(f"Last data row: {last_data_row}.")
    print()

    results = []
    rows_added = 0
    passed = 0
    rejected = 0

    for filepath in resume_files:
        print(f"Screening: {filepath.name} ...", end=" ", flush=True)

        text = extract_text(filepath)
        if not text:
            print("SKIPPED (no text extracted)")
            filepath.unlink()
            continue

        is_resume, reason = validate_is_resume(client, text, filepath.name)
        if not is_resume:
            print(f"REJECTED — {filepath.name} does not appear to be a resume: {reason}")
            filepath.unlink()
            print(f"  (deleted from resumes folder)")
            continue

        info = screen_candidate(client, text, jd_text)

        score = info.get("score", 0)
        label = info.get("label", "Rejected")
        reason = info.get("reason", "")
        name = info.get("name", "")
        name = name.title() if name else "Unknown"

        results.append({
            "name": name,
            "score": score,
            "label": label,
            "reason": reason,
        })

        if score >= 6:
            if recruiter:
                info["_recruiter"] = recruiter
            if client_name:
                info["_client"] = client_name
            if manager_name:
                info["_manager"] = manager_name
            if requirement:
                info["_requirement"] = requirement
            row_data = build_row(info)
            match_score = f"{score}/10 — {label} — {reason}"
            append_row_to_sheet(spreadsheet, ws, row_data, match_score)
            rows_added += 1
            passed += 1
            print(f"PASSED — {name} ({score}/10 {label})")
        else:
            rejected += 1
            print(f"REJECTED — {name} ({score}/10 {label})")

        filepath.unlink()
        print(f"  (deleted from resumes folder)")

    # Print ranked summary
    results.sort(key=lambda r: r["score"], reverse=True)

    print()
    print("=" * 60)
    print("SCREENING RESULTS (ranked by score)")
    print("=" * 60)
    for i, r in enumerate(results, 1):
        status = "PASS" if r["score"] >= 6 else "REJECT"
        print(f"  {i}. {r['name']:30s} {r['score']:>2}/10  {r['label']:20s} [{status}]")
        print(f"     {r['reason']}")
    print("-" * 60)
    print(f"  Passed   : {passed}")
    print(f"  Rejected : {rejected}")
    print(f"  Rows added to CRM : {rows_added}")
    print("=" * 60)


def main():
    parser = argparse.ArgumentParser(description="Agent — Resume Parser & Screener")
    parser.add_argument(
        "--screen", action="store_true",
        help="Enable screening mode: score resumes against jd.txt",
    )
    parser.add_argument(
        "--screen-profile", action="store_true",
        help="Screen a single resume against a JD and save to Screened Profile Tracker",
    )
    parser.add_argument("--resume-file", type=str, default="", help="Path to single resume file")
    parser.add_argument("--jd-file", type=str, default="", help="Path to JD file")
    parser.add_argument("--recruiter", type=str, default="", help="Recruiter name")
    parser.add_argument("--client", type=str, default="", help="Client name")
    parser.add_argument("--manager", type=str, default="", help="Manager name")
    parser.add_argument("--requirement", type=str, default="", help="Requirement/role")
    args = parser.parse_args()

    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        print("Error: ANTHROPIC_API_KEY environment variable is not set.")
        sys.exit(1)

    client = Anthropic(api_key=api_key)

    # Screen Profile mode — standalone, does not use CRM sheet
    if args.screen_profile:
        if not args.resume_file or not args.jd_file or not args.recruiter:
            print("Error: --screen-profile requires --resume-file, --jd-file, and --recruiter")
            sys.exit(1)
        print("Mode: SCREEN PROFILE (single resume screening)")
        run_screen_profile_mode(client, args.resume_file, args.jd_file, args.recruiter)
        return

    print("Connecting to Google Sheet...")
    spreadsheet, ws = connect_to_sheet()
    print(f"Connected to sheet: '{SHEET_NAME}'")

    resume_files = sorted(
        f for f in RESUMES_DIR.iterdir()
        if f.is_file() and f.suffix.lower() in (".pdf", ".docx", ".doc")
    )

    if not resume_files:
        print(f"No PDF or DOCX resumes found in {RESUMES_DIR}/")
        sys.exit(0)

    print(f"Found {len(resume_files)} resume(s) in {RESUMES_DIR}/")

    if args.screen:
        print("Mode: SCREENING (scoring against JD)")
        run_screen_mode(client, resume_files, spreadsheet, ws, recruiter=args.recruiter,
                         client_name=args.client, manager_name=args.manager, requirement=args.requirement)
    else:
        print("Mode: PARSE (extract and append)")
        run_parse_mode(client, resume_files, spreadsheet, ws, recruiter=args.recruiter,
                        client_name=args.client, manager_name=args.manager, requirement=args.requirement)


if __name__ == "__main__":
    main()
